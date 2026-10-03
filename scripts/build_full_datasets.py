import os
import re
import json
import glob
import openpyxl
import numpy as np

# -------------------------------------------------------------------------
# 1. Definisi 35 Kabupaten/Kota di Jawa Tengah
# -------------------------------------------------------------------------
JATENG_KABKOT = [
    {"id": "3301", "name": "Cilacap", "type": "Kabupaten", "karesidenan": "Banyumas"},
    {"id": "3302", "name": "Banyumas", "type": "Kabupaten", "karesidenan": "Banyumas"},
    {"id": "3303", "name": "Purbalingga", "type": "Kabupaten", "karesidenan": "Banyumas"},
    {"id": "3304", "name": "Banjarnegara", "type": "Kabupaten", "karesidenan": "Banyumas"},
    {"id": "3305", "name": "Kebumen", "type": "Kabupaten", "karesidenan": "Kedu"},
    {"id": "3306", "name": "Purworejo", "type": "Kabupaten", "karesidenan": "Kedu"},
    {"id": "3307", "name": "Wonosobo", "type": "Kabupaten", "karesidenan": "Kedu"},
    {"id": "3308", "name": "Magelang", "type": "Kabupaten", "karesidenan": "Kedu"},
    {"id": "3309", "name": "Boyolali", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3310", "name": "Klaten", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3311", "name": "Sukoharjo", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3312", "name": "Wonogiri", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3313", "name": "Karanganyar", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3314", "name": "Sragen", "type": "Kabupaten", "karesidenan": "Surakarta"},
    {"id": "3315", "name": "Grobogan", "type": "Kabupaten", "karesidenan": "Semarang"},
    {"id": "3316", "name": "Blora", "type": "Kabupaten", "karesidenan": "Pati"},
    {"id": "3317", "name": "Rembang", "type": "Kabupaten", "karesidenan": "Pati"},
    {"id": "3318", "name": "Pati", "type": "Kabupaten", "karesidenan": "Pati"},
    {"id": "3319", "name": "Kudus", "type": "Kabupaten", "karesidenan": "Pati"},
    {"id": "3320", "name": "Jepara", "type": "Kabupaten", "karesidenan": "Pati"},
    {"id": "3321", "name": "Demak", "type": "Kabupaten", "karesidenan": "Semarang"},
    {"id": "3322", "name": "Semarang", "type": "Kabupaten", "karesidenan": "Semarang"},
    {"id": "3323", "name": "Temanggung", "type": "Kabupaten", "karesidenan": "Kedu"},
    {"id": "3324", "name": "Kendal", "type": "Kabupaten", "karesidenan": "Semarang"},
    {"id": "3325", "name": "Batang", "type": "Kabupaten", "karesidenan": "Pekalongan"},
    {"id": "3326", "name": "Pekalongan", "type": "Kabupaten", "karesidenan": "Pekalongan"},
    {"id": "3327", "name": "Pemalang", "type": "Kabupaten", "karesidenan": "Pekalongan"},
    {"id": "3328", "name": "Tegal", "type": "Kabupaten", "karesidenan": "Pekalongan"},
    {"id": "3329", "name": "Brebes", "type": "Kabupaten", "karesidenan": "Pekalongan"},
    {"id": "3371", "name": "Kota Magelang", "type": "Kota", "karesidenan": "Kedu"},
    {"id": "3372", "name": "Kota Surakarta", "type": "Kota", "karesidenan": "Surakarta"},
    {"id": "3373", "name": "Kota Salatiga", "type": "Kota", "karesidenan": "Semarang"},
    {"id": "3374", "name": "Kota Semarang", "type": "Kota", "karesidenan": "Semarang"},
    {"id": "3375", "name": "Kota Pekalongan", "type": "Kota", "karesidenan": "Pekalongan"},
    {"id": "3376", "name": "Kota Tegal", "type": "Kota", "karesidenan": "Pekalongan"}
]

JATENG_CODES = {item["id"]: item["name"] for item in JATENG_KABKOT}

def resolve_jateng_id(name_str):
    name_str = str(name_str).strip()
    m = re.match(r'^(33\d{2})', name_str)
    if m and m.group(1) in JATENG_CODES:
        return m.group(1)
    clean = re.sub(r'^(kabupaten|kab\.)\s*', '', name_str, flags=re.I).strip()
    if clean.lower().startswith('kota '):
        for code, name in JATENG_CODES.items():
            if name.lower() == clean.lower():
                return code
    else:
        for code, name in JATENG_CODES.items():
            if not name.lower().startswith('kota ') and name.lower() == clean.lower():
                return code
    return None

def parse_num(val):
    if val is None:
        return None
    s = str(val).strip().replace(',', '.')
    if s in ['...', '-', 'null', 'None', '']:
        return None
    try:
        return float(s)
    except:
        return None

# -------------------------------------------------------------------------
# 2. Ekstraksi Data Multivariat 2021-2025
# -------------------------------------------------------------------------
print("--> Memulai ekstraksi data Multivariat Jawa Tengah 2021-2025...")
years = [2021, 2022, 2023, 2024, 2025]

# Data dictionary: year -> cid -> var_name -> value
mv_data = {yr: {item["id"]: {
    "id": item["id"],
    "name": item["name"],
    "type": item["type"],
    "karesidenan": item["karesidenan"],
    "year": yr
} for item in JATENG_KABKOT} for yr in years}

# A. PDRB per Kapita (Ribu Rp)
for yr in years:
    f = glob.glob(f'Data PDRB per Kapita 2021-2025/*{yr}*.xlsx')[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, 2).value)
            mv_data[yr][cid]["pdrb_kapita"] = val

# B. TPT & TPAK (Agustus)
for yr in years:
    f = glob.glob(f'Data TPT & TPAK Agustus 2021-2025/*{yr}*.xlsx')[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            # Col 3 = TPT Agustus, Col 5 = TPAK Agustus
            tpt_val = parse_num(ws.cell(r, 3).value)
            tpak_val = parse_num(ws.cell(r, 5).value)
            mv_data[yr][cid]["tpt"] = tpt_val
            mv_data[yr][cid]["tpak"] = tpak_val

# C. UMK (Rp)
for yr in years:
    f = glob.glob(f'Data UMK 2021-2025/*{yr}*.xlsx')[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, 2).value)
            mv_data[yr][cid]["umk"] = val

# D. Kepadatan Penduduk (jiwa/km2)
for yr in years:
    f = glob.glob(f'Kepadatan Penduduk 2021-2025/*{yr}*.xlsx')[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, 2).value)
            mv_data[yr][cid]["kepadatan"] = val

# E. Rata-Rata Lama Sekolah (tahun)
for yr in years:
    f = glob.glob(f'Rata-Rata Lama Sekolah 2021-2025/*{yr}*.xlsx')[0]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, 2).value)
            mv_data[yr][cid]["rls"] = val

# F. Rata-Rata Pengeluaran per Kapita (Rp/bulan)
for yr in years:
    matches = [m for m in glob.glob(f'Rata-Rata Pengeluaran Perkapita/*{yr}*.xlsx') if '~$' not in m]
    wb = openpyxl.load_workbook(matches[0], data_only=True)
    ws = wb.active
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, 2).value)
            mv_data[yr][cid]["pengeluaran"] = val

# G. Total Angkatan Kerja (jiwa)
for yr in years:
    matches = [m for m in glob.glob(f'Total Angkatan Kerja 2021-2025/*{yr}*.xlsx') if '~$' not in m]
    wb = openpyxl.load_workbook(matches[0], data_only=True)
    ws = wb.active
    col_idx = 2 if yr == 2024 else 4  # Tahun 2024 kolom 2, tahun lainnya kolom 4
    for r in range(2, ws.max_row+1):
        cid = resolve_jateng_id(ws.cell(r, 1).value)
        if cid:
            val = parse_num(ws.cell(r, col_idx).value)
            mv_data[yr][cid]["angkatan_kerja"] = val

# Verifikasi kelengkapan
var_keys = ["tpt", "tpak", "umk", "pdrb_kapita", "pengeluaran", "angkatan_kerja", "rls", "kepadatan"]
all_complete = True
for yr in years:
    for cid in JATENG_CODES:
        for vk in var_keys:
            if vk not in mv_data[yr][cid] or mv_data[yr][cid][vk] is None:
                print(f"ERROR: missing {vk} in year {yr} for {cid} ({JATENG_CODES[cid]})")
                all_complete = False

if all_complete:
    print("[OK] SELURUH DATA MULTIVARIAT (8 variabel x 35 kab/kota x 5 tahun) 100% LENGKAP TANPA KOSONG!")

# -------------------------------------------------------------------------
# 3. Hitung PCA (Principal Component Analysis)
# -------------------------------------------------------------------------
print("--> Menghitung PCA untuk data multivariat...")

# Kita hitung PCA untuk tahun 2025 (tahun utama) dan juga memproyeksikan seluruh tahun
# Siapkan matriks X untuk 2025 (35 baris x 8 kolom)
cids = sorted(list(JATENG_CODES.keys()))
X_2025 = []
for cid in cids:
    row = [mv_data[2025][cid][vk] for vk in var_keys]
    X_2025.append(row)
X_2025 = np.array(X_2025, dtype=float)

# Standardisasi Z = (X - mean) / std
mu = np.mean(X_2025, axis=0)
sigma = np.std(X_2025, axis=0, ddof=1)
Z_2025 = (X_2025 - mu) / sigma

# SVD
U, S, Vt = np.linalg.svd(Z_2025, full_matrices=False)
eigenvalues = (S ** 2) / (len(cids) - 1)
total_var = np.sum(eigenvalues)
var_explained = eigenvalues / total_var
cum_var_explained = np.cumsum(var_explained)

# Loadings: korelasi antara variabel dengan PC = Vt.T * sqrt(eigenvalues)
# atau langsung arah komponen utama (eigenvectors)
loadings_matrix = Vt.T  # 8 x 8

pca_meta = {
    "variables": var_keys,
    "labels": [
        "Tingkat Pengangguran Terbuka (%)",
        "Tingkat Partisipasi Angkatan Kerja (%)",
        "Upah Minimum Kab/Kota (Rp)",
        "PDRB per Kapita (Ribu Rp)",
        "Pengeluaran per Kapita (Rp/bln)",
        "Jumlah Angkatan Kerja (jiwa)",
        "Rata-Rata Lama Sekolah (thn)",
        "Kepadatan Penduduk (jiwa/km²)"
    ],
    "eigenvalues": [round(float(e), 4) for e in eigenvalues[:4]],
    "variance_explained": [round(float(v * 100), 2) for v in var_explained[:4]],
    "cum_variance_explained": [round(float(c * 100), 2) for c in cum_var_explained[:4]],
    "loadings": {
        var_keys[i]: {
            "PC1": round(float(loadings_matrix[i, 0]), 4),
            "PC2": round(float(loadings_matrix[i, 1]), 4),
            "PC3": round(float(loadings_matrix[i, 2]), 4)
        } for i in range(len(var_keys))
    }
}

print(f"PCA Variance Explained: PC1 = {pca_meta['variance_explained'][0]}%, PC2 = {pca_meta['variance_explained'][1]}% (Total 2D: {pca_meta['cum_variance_explained'][1]}%)")

# Proyeksikan scores untuk semua tahun menggunakan basis loading 2025 agar konsisten
for yr in years:
    X_yr = np.array([[mv_data[yr][cid][vk] for vk in var_keys] for cid in cids], dtype=float)
    Z_yr = (X_yr - mu) / sigma
    scores = np.dot(Z_yr, loadings_matrix)
    for idx, cid in enumerate(cids):
        mv_data[yr][cid]["pca"] = {
            "pc1": round(float(scores[idx, 0]), 4),
            "pc2": round(float(scores[idx, 1]), 4),
            "pc3": round(float(scores[idx, 2]), 4)
        }
        # Tentukan zona industri / klaster analitis
        # Batang, Kendal, Brebes -> Kawasan Industri Baru Pantura (relokasi pabrik)
        if cid in ["3325", "3324", "3329"]:
            mv_data[yr][cid]["cluster"] = "Kawasan Industri Baru (Relokasi)"
        elif cid in ["3374", "3322", "3372", "3311"]:
            mv_data[yr][cid]["cluster"] = "Pusat Metropolis & Industri Matang"
        elif mv_data[yr][cid]["type"] == "Kota":
            mv_data[yr][cid]["cluster"] = "Kota Jasa & Perdagangan"
        elif cid in ["3301", "3319", "3320"]:
            mv_data[yr][cid]["cluster"] = "Sentra Industri Spesifik / Energi"
        else:
            mv_data[yr][cid]["cluster"] = "Wilayah Agraris & Penyangga"

# Simpan data multivariat
os.makedirs("data", exist_ok=True)
with open("data/multivariate_jateng.json", "w", encoding="utf-8") as f:
    json.dump({
        "metadata": {
            "source": "BPS Provinsi Jawa Tengah (2021-2025)",
            "unit_count": len(cids),
            "years": years,
            "variables": var_keys,
            "pca": pca_meta
        },
        "records": mv_data
    }, f, indent=2)
print("[OK] Berhasil menyimpan data/multivariate_jateng.json")

# -------------------------------------------------------------------------
# 4. Data Hierarki Jawa Tengah (Eks-Karesidenan -> Kab/Kota -> Sektor)
# -------------------------------------------------------------------------
print("--> Membangun struktur hierarki Jawa Tengah...")
# Ambil data sektor dari 'Data Geospatial Seluruh Kabupaten Kota 2025/Penduduk Kerja 2025 xlsx.xlsx'
f_pend_kerja = 'Data Geospatial Seluruh Kabupaten Kota 2025/Penduduk Kerja 2025 xlsx.xlsx'
wb_pk = openpyxl.load_workbook(f_pend_kerja, data_only=True)
ws_pk = wb_pk.active
sektor_data = {}
for r in range(4, ws_pk.max_row+1):
    c1 = ws_pk.cell(r, 1).value
    if not c1: continue
    cid = resolve_jateng_id(c1)
    if cid:
        pertanian = parse_num(ws_pk.cell(r, 2).value) or 0
        industri = parse_num(ws_pk.cell(r, 3).value) or 0
        jasa = parse_num(ws_pk.cell(r, 4).value) or 0
        total = pertanian + industri + jasa
        sektor_data[cid] = {
            "Pertanian": pertanian,
            "Industri Pengolahan": industri,
            "Jasa & Lainnya": jasa,
            "Total": total
        }

# Bentuk tree structure: Root -> Karesidenan -> Kab/Kota -> Sektor
hierarchy_root = {
    "name": "Provinsi Jawa Tengah",
    "level": "Provinsi",
    "children": []
}

karesidenan_map = {}
for item in JATENG_KABKOT:
    k_name = f"Karesidenan {item['karesidenan']}"
    if k_name not in karesidenan_map:
        k_node = {
            "name": k_name,
            "karesidenan": item["karesidenan"],
            "level": "Karesidenan",
            "children": []
        }
        karesidenan_map[k_name] = k_node
        hierarchy_root["children"].append(k_node)
    
    cid = item["id"]
    kab_info = mv_data[2025][cid]
    sektor_info = sektor_data.get(cid, {"Pertanian": 0, "Industri Pengolahan": 0, "Jasa & Lainnya": 0})
    
    kab_node = {
        "id": cid,
        "name": item["name"],
        "type": item["type"],
        "karesidenan": item["karesidenan"],
        "level": "Kabupaten/Kota",
        "umk": kab_info["umk"],
        "tpt": kab_info["tpt"],
        "pdrb_kapita": kab_info["pdrb_kapita"],
        "pengeluaran": kab_info["pengeluaran"],
        "cluster": kab_info["cluster"],
        "children": []
    }
    
    # Tambahkan sektor sebagai daun (leaves)
    for sektor_name, sektor_val in [
        ("Industri Pengolahan", sektor_info["Industri Pengolahan"]),
        ("Pertanian, Kehutanan & Perikanan", sektor_info["Pertanian"]),
        ("Perdagangan & Jasa", sektor_info["Jasa & Lainnya"])
    ]:
        kab_node["children"].append({
            "name": sektor_name,
            "level": "Sektor",
            "kabupaten": item["name"],
            "karesidenan": item["karesidenan"],
            "value": int(sektor_val),       # Size = Jumlah Tenaga Kerja
            "umk": kab_info["umk"],          # Color metric 1 = UMK
            "pengeluaran": kab_info["pengeluaran"], # Color metric 2 = Pengeluaran
            "tpt": kab_info["tpt"]
        })
    
    karesidenan_map[k_name]["children"].append(kab_node)

with open("data/hierarchy_jateng.json", "w", encoding="utf-8") as f:
    json.dump(hierarchy_root, f, indent=2)
print("[OK] Berhasil menyimpan data/hierarchy_jateng.json")

# -------------------------------------------------------------------------
# 5. Data Geospasial Nasional (515 Kab/Kota LapakGIS 2024 + BPS 2025)
# -------------------------------------------------------------------------
print("--> Memproses data geospasial nasional menggunakan Shapefile LapakGIS 2024...")
import sys
if os.path.dirname(__file__) not in sys.path:
    sys.path.insert(0, os.path.dirname(__file__))
from process_lapakgis_shapefile import run as process_shp
process_shp()
print("=== SELURUH DATA PIPELINE SELESAI DENGAN SUKSES! ===")

